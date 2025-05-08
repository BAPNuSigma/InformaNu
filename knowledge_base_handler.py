import os
import markdown
from pathlib import Path
from datetime import datetime
import re
from docx import Document
import openpyxl

class KnowledgeBaseHandler:
    def __init__(self, knowledge_base_dir="knowledge_base"):
        self.knowledge_base_dir = knowledge_base_dir
        self.knowledge_base = {}
        self.load_knowledge_base()
    
    def load_knowledge_base(self):
        """Load all supported files from the knowledge base directory"""
        if not os.path.exists(self.knowledge_base_dir):
            raise FileNotFoundError(f"Knowledge base directory {self.knowledge_base_dir} not found")
        
        for file_path in Path(self.knowledge_base_dir).glob("*"):
            if file_path.suffix == ".md":
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                    self.knowledge_base[file_path.stem] = {
                        'markdown': content,
                        'text': markdown.markdown(content)
                    }
            elif file_path.suffix == ".docx":
                doc = Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                content = '\n'.join(paragraphs)
                self.knowledge_base[file_path.stem] = {
                    'markdown': content,
                    'text': content
                }
            elif file_path.suffix == ".xlsx":
                wb = openpyxl.load_workbook(file_path)
                all_text = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = [str(cell) for cell in row if cell is not None]
                        if row_text:
                            all_text.append(' | '.join(row_text))
                content = '\n'.join(all_text)
                self.knowledge_base[file_path.stem] = {
                    'markdown': content,
                    'text': content
                }

    def extract_section(self, content, section_header):
        """Extract a specific section from the content"""
        lines = content.split('\n')
        section_content = []
        in_section = False
        
        for line in lines:
            if line.strip().lower().startswith(section_header.lower()):
                in_section = True
                section_content.append(line.strip())
            elif in_section:
                if line.strip() and not line.startswith('#'):
                    section_content.append(line.strip())
                elif line.startswith('#'):
                    break
        
        return '\n'.join(section_content) if section_content else None

    def get_next_meeting(self):
        """Find the next meeting from the schedule based on current date"""
        schedule = self.knowledge_base.get('spring_2025_schedule', {}).get('markdown', '')
        if not schedule:
            return "I couldn't find the meeting schedule."
        
        current_date = datetime.now()
        next_meeting = None
        next_meeting_date = None
        next_meeting_details = []

        # Parse the schedule looking for meeting dates
        lines = schedule.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            # Look for date patterns like "1/22/25" or similar
            date_match = re.search(r'(\d{1,2}/\d{1,2}/\d{2,4})', line)
            time_match = re.search(r'(\d{1,2}:\d{2}(?:\s*[AaPp][Mm])?)', line)
            
            if date_match and time_match:
                date_str = date_match.group(1)
                time_str = time_match.group(1)
                
                try:
                    # Ensure 4-digit year
                    if len(date_str.split('/')[-1]) == 2:
                        date_str = date_str[:-2] + '20' + date_str[-2:]
                    
                    # Standardize time format
                    if 'PM' in time_str.upper() and not time_str.startswith('12'):
                        hour = int(time_str.split(':')[0]) + 12
                        time_str = f"{hour}:{time_str.split(':')[1]}"
                    time_str = time_str.replace('AM', '').replace('PM', '').strip()
                    
                    meeting_datetime = datetime.strptime(f"{date_str} {time_str}", '%m/%d/%Y %H:%M')
                    
                    if meeting_datetime > current_date and (next_meeting_date is None or meeting_datetime < next_meeting_date):
                        next_meeting_details = [line]
                        # Gather additional details from subsequent lines
                        j = i + 1
                        while j < len(lines) and (lines[j].strip().startswith('-') or lines[j].strip().startswith('  ')):
                            next_meeting_details.append(lines[j].strip())
                            j += 1
                        next_meeting_date = meeting_datetime
                except ValueError:
                    pass
            i += 1

        if next_meeting_details:
            # Format the meeting information nicely
            return "\n".join(next_meeting_details)
        else:
            return "I couldn't find any upcoming meetings in the schedule."

    def get_gpa_requirement(self):
        """Get the specific GPA requirement for our chapter"""
        requirements = self.knowledge_base.get('membership_types_and_requirements', {}).get('markdown', '')
        if not requirements:
            return "I couldn't find the GPA requirements."
        
        gpa_section = self.extract_section(requirements, "GPA Requirements:")
        if gpa_section:
            return gpa_section
        return "I couldn't find the specific GPA requirements."

    def get_attendance_requirements(self):
        """Get the specific attendance requirements"""
        requirements = self.knowledge_base.get('member_requirements', {}).get('markdown', '')
        if not requirements:
            return "I couldn't find the attendance requirements."
        
        attendance_section = self.extract_section(requirements, "Attendance Requirements")
        if attendance_section:
            return attendance_section
        return "I couldn't find the specific attendance requirements."

    def get_membership_checklist(self):
        """Get the specific membership checklist"""
        requirements = self.knowledge_base.get('member_requirements', {}).get('markdown', '')
        if not requirements:
            return "I couldn't find the membership checklist."
        
        checklist_section = self.extract_section(requirements, "Member Checklist")
        if checklist_section:
            return checklist_section
        return "I couldn't find the specific membership checklist."

    def get_reflection_papers_policy(self):
        """Get the specific reflection papers policy"""
        requirements = self.knowledge_base.get('membership_types_and_requirements', {}).get('markdown', '')
        if not requirements:
            return "I couldn't find the reflection papers policy."
        
        policy_section = self.extract_section(requirements, "Reflection Papers Policy")
        if policy_section:
            return policy_section
        return "I couldn't find the specific reflection papers policy."

    def get_tutoring_schedule(self):
        """Get the specific tutoring schedule"""
        requirements = self.knowledge_base.get('membership_types_and_requirements', {}).get('markdown', '')
        if not requirements:
            return "I couldn't find the tutoring schedule."
        
        schedule_section = self.extract_section(requirements, "Tutoring Schedule")
        if schedule_section:
            return schedule_section
        return "I couldn't find the specific tutoring schedule."

    def search_knowledge_base(self, query):
        """Enhanced search implementation with specific handlers for common queries"""
        query = query.lower()
        
        # Special handling for common queries
        if 'next meeting' in query or ('when' in query and 'meeting' in query):
            return [{'content': self.get_next_meeting(), 'relevance': 1.0}]
        elif 'gpa' in query and ('requirement' in query or 'need' in query or 'join' in query):
            return [{'content': self.get_gpa_requirement(), 'relevance': 1.0}]
        elif 'attendance' in query and 'requirement' in query:
            return [{'content': self.get_attendance_requirements(), 'relevance': 1.0}]
        elif 'checklist' in query:
            return [{'content': self.get_membership_checklist(), 'relevance': 1.0}]
        elif 'reflection' in query and 'paper' in query:
            return [{'content': self.get_reflection_papers_policy(), 'relevance': 1.0}]
        elif 'tutoring' in query and 'schedule' in query:
            return [{'content': self.get_tutoring_schedule(), 'relevance': 1.0}]
        
        # Default search behavior for other queries
        matches = []
        for topic, content in self.knowledge_base.items():
            text = content['text'].lower()
            if query in text:
                # Find the most relevant section
                best_section = None
                best_relevance = 0
                
                sections = text.split('\n\n')
                for section in sections:
                    if query in section.lower():
                        # Calculate relevance score based on query term frequency and section length
                        relevance = section.lower().count(query) / len(section)
                        if relevance > best_relevance:
                            best_section = section
                            best_relevance = relevance
                
                if best_section:
                    matches.append({
                        'topic': topic,
                        'content': best_section,
                        'relevance': best_relevance
                    })
        
        matches.sort(key=lambda x: x['relevance'], reverse=True)
        return matches

    def get_topic_content(self, topic):
        """Get the full content for a specific topic"""
        return self.knowledge_base.get(topic, {}).get('markdown', '') 